"""Turn an uploaded roster file into structured weeks.

Shops keep their history in whatever they already use — a spreadsheet, a
scanned PDF, or a photo of the sheet pinned to the staff-room wall. All three
have to work, because "export your data as CSV first" is exactly the barrier
that stops a new customer ever getting started.

Routing by file type:

    .xlsx .xlsm     openpyxl        -> structured, exact
    .csv  .tsv      csv module      -> structured, exact
    .pdf            rendered to images, then read by the model
    .png .jpg ...   read by the model

Spreadsheets are parsed deterministically and are always preferred. Images
and PDFs go through the LLM, which is slower, costs money, and is
occasionally wrong — so anything it produces is presented for review rather
than imported directly.
"""
from __future__ import annotations

import csv
import io
import logging
from datetime import datetime, timedelta
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

from app.services import llm
from app.services.roster_import import ParsedWeek, parse_sheet, parse_workbook_rows

log = logging.getLogger("roster.file_import")

# Generous enough for a year of rosters, small enough that a mis-drag of a
# video file fails fast rather than filling memory.
MAX_UPLOAD_BYTES = 25 * 1024 * 1024

# Each PDF page becomes one model call, so a mistakenly uploaded 200-page
# document would be slow and expensive.
MAX_PDF_PAGES = 30

# 2x scale: roster text is small and renders illegibly at native resolution,
# but going higher inflates the image without improving the read.
PDF_RENDER_SCALE = 2.0

SPREADSHEET_EXTENSIONS = (".xlsx", ".xlsm", ".xltx")
CSV_EXTENSIONS = (".csv", ".tsv", ".txt")
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".heic")
PDF_EXTENSIONS = (".pdf",)

SUPPORTED_EXTENSIONS = (
    SPREADSHEET_EXTENSIONS + CSV_EXTENSIONS + IMAGE_EXTENSIONS + PDF_EXTENSIONS
)


class UnsupportedFile(ValueError):
    """The upload is not a file type we can read."""


@dataclass
class ImportResult:
    weeks: List[ParsedWeek]
    kind: str                       # spreadsheet | csv | image | pdf
    pages: int = 1
    needs_review: bool = False      # model-derived data always does
    notes: List[str] = None

    def __post_init__(self):
        if self.notes is None:
            self.notes = []


def detect_kind(filename: str, content_type: str = "") -> str:
    name = (filename or "").lower()
    if name.endswith(SPREADSHEET_EXTENSIONS):
        return "spreadsheet"
    if name.endswith(CSV_EXTENSIONS):
        return "csv"
    if name.endswith(PDF_EXTENSIONS):
        return "pdf"
    if name.endswith(IMAGE_EXTENSIONS):
        return "image"

    # Fall back to the browser's content type when the name is unhelpful.
    if content_type.startswith("image/"):
        return "image"
    if content_type == "application/pdf":
        return "pdf"
    if "spreadsheet" in content_type or "excel" in content_type:
        return "spreadsheet"
    if content_type in ("text/csv", "text/plain"):
        return "csv"

    raise UnsupportedFile(
        f"Cannot read '{filename}'. Upload a spreadsheet (.xlsx/.csv), "
        f"a PDF, or a photo of the roster."
    )


# ---------------------------------------------------------------------------
# Spreadsheets and CSV — deterministic
# ---------------------------------------------------------------------------
def _read_spreadsheet(data: bytes, year: int) -> Tuple[List[ParsedWeek], int]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets = [
        (name, [list(row) for row in workbook[name].iter_rows(values_only=True)])
        for name in workbook.sheetnames
    ]
    workbook.close()
    return parse_workbook_rows(sheets, year), len(sheets)


def _read_csv(data: bytes, filename: str, year: int) -> List[ParsedWeek]:
    # Roster files are routinely saved from Excel on Windows, so UTF-8 is not
    # a safe assumption — cp1252 and latin-1 both appear in the wild, and
    # latin-1 never raises, which makes it a usable last resort.
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise UnsupportedFile("Could not decode that file as text.")

    delimiter = "\t" if filename.lower().endswith(".tsv") else ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    # A CSV is one sheet, and its name is the only place a date can come from.
    return [parse_sheet(filename.rsplit(".", 1)[0], rows, year)]


# ---------------------------------------------------------------------------
# Images and PDFs — read by the model
# ---------------------------------------------------------------------------
def _pdf_to_images(data: bytes) -> List[bytes]:
    try:
        import pymupdf
    except ImportError:                                    # pragma: no cover
        raise UnsupportedFile(
            "PDF support needs the 'pymupdf' package. Run: pip install -r requirements.txt"
        )

    document = pymupdf.open(stream=data, filetype="pdf")
    if document.page_count > MAX_PDF_PAGES:
        document.close()
        raise UnsupportedFile(
            f"That PDF has {document.page_count} pages; the limit is {MAX_PDF_PAGES}. "
            f"Split it or upload the roster pages only."
        )

    matrix = pymupdf.Matrix(PDF_RENDER_SCALE, PDF_RENDER_SCALE)
    images = [
        document.load_page(number).get_pixmap(matrix=matrix).tobytes("png")
        for number in range(document.page_count)
    ]
    document.close()
    return images


async def _read_images(
    images: List[bytes], known_names: List[str], year: int
) -> Tuple[List[ParsedWeek], List[str]]:
    """OCR each image into a week. One bad page does not lose the others."""
    weeks: List[ParsedWeek] = []
    notes: List[str] = []

    for index, image in enumerate(images, start=1):
        label = f"page {index}" if len(images) > 1 else "image"
        try:
            parsed = await llm.ocr_roster_bytes(image, known_names)
        except llm.LLMUnavailable:
            raise
        except Exception as exc:
            log.warning("OCR failed on %s: %s", label, exc)
            notes.append(f"{label}: could not be read ({str(exc)[:120]})")
            continue

        week = _week_from_ocr(parsed, label, year)
        if week.shifts:
            weeks.append(week)
        else:
            notes.append(f"{label}: no shifts recognised")

    return weeks, notes


def _week_from_ocr(parsed: Dict[str, Any], label: str, year: int) -> ParsedWeek:
    """Reshape the model's JSON into the same structure a spreadsheet gives.

    Everything downstream — preview, review, commit — then works identically
    regardless of where the data came from.
    """
    from app.services.roster_import import (
        ParsedShift, normalise_name, parse_shift_cell, parse_week_ending,
    )

    week = ParsedWeek(
        sheet_name=label,
        week_ending_label=parsed.get("week_start"),
        week_start=None,
    )

    raw_week = parsed.get("week_start")
    if raw_week:
        # The model is asked for the Monday directly, so this is a plain date
        # rather than the week-ending convention used by sheet names.
        try:
            from datetime import datetime, timedelta
            monday = datetime.strptime(raw_week, "%Y-%m-%d").date()
            monday -= timedelta(days=monday.weekday())
            week.week_start = monday.isoformat()
        except (ValueError, TypeError):
            parsed_date = parse_week_ending(str(raw_week), year)
            if parsed_date:
                week.week_start = parsed_date.isoformat()

    if not week.week_start:
        week.warnings.append(
            f"{label}: could not determine which week this is — set it before importing."
        )

    for entry in parsed.get("shifts", []):
        name = normalise_name(entry.get("employee_name"))
        day = (entry.get("day") or "").strip().lower()[:3]
        if not name or day not in ("mon", "tue", "wed", "thu", "fri", "sat", "sun"):
            continue

        start, end = entry.get("start"), entry.get("end")
        if not (start and end):
            continue

        shift = parse_shift_cell(f"{start}-{end}")
        if not shift:
            week.unparsed_cells.append({
                "employee_name": name, "day": day,
                "value": f"{start}-{end}", "reason": "unreadable times",
            })
            continue

        week.employees.setdefault(name, entry.get("role"))
        week.shifts.append(ParsedShift(
            employee_name=name,
            role=entry.get("role"),
            day=day,
            start=shift["start"],
            end=shift["end"],
            hours=shift["hours"],
            overnight=shift["overnight"],
        ))

    return week


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def _apply_week_start_fallback(
    weeks: List[ParsedWeek], week_start: Optional[str]
) -> List[ParsedWeek]:
    """Date any weeks whose own file didn't say when they were.

    A multi-sheet workbook names its weeks ('we Aug 23rd 26'), but a CSV
    export or a photo usually does not — the file is just "this week's rota".
    Without a date those weeks are unusable and the whole upload gets
    rejected despite having parsed perfectly, so the caller can supply the
    first week and later undated weeks follow on seven days at a time.
    """
    if not week_start:
        return weeks
    try:
        cursor = datetime.strptime(week_start, "%Y-%m-%d").date()
    except ValueError:
        return weeks
    # Snap to Monday: every week in this system starts there.
    cursor -= timedelta(days=cursor.weekday())

    for week in weeks:
        if week.week_start or not week.shifts:
            continue
        week.week_start = cursor.isoformat()
        week.warnings.append(
            f"No date in the file — assigned to the week beginning {cursor.isoformat()}."
        )
        cursor += timedelta(days=7)
    return weeks


async def parse_upload(
    filename: str,
    data: bytes,
    *,
    content_type: str = "",
    known_names: Optional[List[str]] = None,
    year: int = 2026,
    week_start: Optional[str] = None,
) -> ImportResult:
    """Parse any supported roster file into weeks ready for review.

    `week_start` dates any week the file itself doesn't identify — needed for
    CSV exports and photos, which rarely carry the week in the filename.
    """
    if not data:
        raise UnsupportedFile("That file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedFile(
            f"That file is {len(data) // 1024 // 1024}MB; the limit is "
            f"{MAX_UPLOAD_BYTES // 1024 // 1024}MB."
        )

    kind = detect_kind(filename, content_type)
    known_names = known_names or []

    if kind == "spreadsheet":
        weeks, sheets = _read_spreadsheet(data, year)
        return ImportResult(
            weeks=_apply_week_start_fallback(weeks, week_start), kind=kind, pages=sheets
        )

    if kind == "csv":
        weeks = _read_csv(data, filename, year)
        return ImportResult(weeks=_apply_week_start_fallback(weeks, week_start), kind=kind)

    images = _pdf_to_images(data) if kind == "pdf" else [data]
    weeks, notes = await _read_images(images, known_names, year)
    weeks = _apply_week_start_fallback(weeks, week_start)
    return ImportResult(
        weeks=weeks,
        kind=kind,
        pages=len(images),
        # Model output is a best-effort reading of a picture, so it is always
        # flagged for a human to check before anything is written.
        needs_review=True,
        notes=notes,
    )
